"""
将 doc/ 目录下所有 .docx 文件转换为 .md 格式
支持：标题层级、段落、有序/无序列表、表格、加粗、斜体
"""

import os
import re
from docx import Document
from docx.oxml.ns import qn

DOC_DIR = os.path.dirname(os.path.abspath(__file__))

HEADING_MAP = {
    'Heading 1': '#',
    'Heading 2': '##',
    'Heading 3': '###',
    'Heading 4': '####',
    'Heading 5': '#####',
    '标题 1': '#',
    '标题 2': '##',
    '标题 3': '###',
    '标题 4': '####',
    'Title': '#',
    '标题': '#',
}


def get_run_text(run):
    """提取 run 的文字，加上加粗/斜体 markdown 符号"""
    text = run.text
    if not text.strip():
        return text
    if run.bold and run.italic:
        return f'***{text}***'
    if run.bold:
        return f'**{text}**'
    if run.italic:
        return f'*{text}*'
    return text


def get_paragraph_text(para):
    """拼接段落内所有 run 的文字"""
    return ''.join(get_run_text(r) for r in para.runs)


def is_list_paragraph(para):
    """判断是否为列表段落"""
    style_name = para.style.name if para.style else ''
    if 'List' in style_name or '列表' in style_name:
        return True
    # 通过 XML 判断 numPr
    pPr = para._p.find(qn('w:pPr'))
    if pPr is not None and pPr.find(qn('w:numPr')) is not None:
        return True
    return False


def get_list_level(para):
    """获取列表缩进层级（0-based）"""
    try:
        pPr = para._p.find(qn('w:pPr'))
        if pPr is not None:
            numPr = pPr.find(qn('w:numPr'))
            if numPr is not None:
                ilvl = numPr.find(qn('w:ilvl'))
                if ilvl is not None:
                    return int(ilvl.get(qn('w:val'), 0))
    except Exception:
        pass
    return 0


def is_ordered_list(para):
    """简单判断是否为有序列表（通过样式名）"""
    style_name = para.style.name if para.style else ''
    return 'Number' in style_name or '编号' in style_name or '有序' in style_name


def table_to_md(table):
    """将 docx 表格转换为 markdown 表格，处理合并单元格去重"""
    lines = []
    for i, row in enumerate(table.rows):
        cells = []
        seen_tc = set()  # 用于检测合并单元格（相同 _tc 对象）
        for cell in row.cells:
            tc_id = id(cell._tc)
            if tc_id in seen_tc:
                # 跳过合并单元格的重复部分
                continue
            seen_tc.add(tc_id)
            cell_text = ' '.join(cell.text.split())
            cells.append(cell_text)
        if not cells:
            continue
        line = '| ' + ' | '.join(cells) + ' |'
        lines.append(line)
        if i == 0:
            sep = '| ' + ' | '.join(['---'] * len(cells)) + ' |'
            lines.append(sep)
    return '\n'.join(lines)


def docx_to_md(docx_path):
    """主转换函数，返回 markdown 字符串"""
    doc = Document(docx_path)
    md_lines = []

    # 用于追踪有序列表计数
    ordered_counters = {}

    for block in doc.element.body:
        tag = block.tag.split('}')[-1] if '}' in block.tag else block.tag

        # 表格
        if tag == 'tbl':
            # 找到对应的 Table 对象
            for table in doc.tables:
                if table._tbl is block:
                    md_lines.append('')
                    md_lines.append(table_to_md(table))
                    md_lines.append('')
                    break
            continue

        # 段落
        if tag == 'p':
            # 找对应的 Paragraph 对象
            para = None
            for p in doc.paragraphs:
                if p._p is block:
                    para = p
                    break
            if para is None:
                continue

            style_name = para.style.name if para.style else 'Normal'
            text = get_paragraph_text(para).strip()

            # 跳过完全空段落（保留一个空行）
            if not text:
                if md_lines and md_lines[-1] != '':
                    md_lines.append('')
                continue

            # 标题
            if style_name in HEADING_MAP:
                prefix = HEADING_MAP[style_name]
                md_lines.append('')
                md_lines.append(f'{prefix} {text}')
                md_lines.append('')
                ordered_counters.clear()
                continue

            # 列表
            if is_list_paragraph(para):
                level = get_list_level(para)
                indent = '  ' * level
                if is_ordered_list(para):
                    counter = ordered_counters.get(level, 0) + 1
                    ordered_counters[level] = counter
                    # 重置更深层级的计数
                    for k in list(ordered_counters.keys()):
                        if k > level:
                            del ordered_counters[k]
                    md_lines.append(f'{indent}{counter}. {text}')
                else:
                    ordered_counters.clear()
                    md_lines.append(f'{indent}- {text}')
                continue

            # 普通段落
            ordered_counters.clear()
            md_lines.append(text)
            md_lines.append('')

    # 清理多余的空行（连续3个以上空行 → 2个）
    result = '\n'.join(md_lines)
    result = re.sub(r'\n{3,}', '\n\n', result)
    return result.strip()


def convert_all():
    docx_files = [
        f for f in os.listdir(DOC_DIR)
        if f.endswith('.docx') and os.path.isfile(os.path.join(DOC_DIR, f))
    ]

    if not docx_files:
        print('doc/ 目录下未找到 .docx 文件')
        return

    success, failed = [], []

    for fname in sorted(docx_files):
        src = os.path.join(DOC_DIR, fname)
        dst = os.path.join(DOC_DIR, fname.replace('.docx', '.md'))
        try:
            md_content = docx_to_md(src)
            with open(dst, 'w', encoding='utf-8') as f:
                f.write(md_content)
            size = len(md_content)
            print(f'  ✓  {fname}  →  {os.path.basename(dst)}  ({size:,} chars)')
            success.append(fname)
        except Exception as e:
            print(f'  ✗  {fname}  错误: {e}')
            failed.append((fname, str(e)))

    print(f'\n完成：{len(success)} 个成功，{len(failed)} 个失败')
    if failed:
        for f, e in failed:
            print(f'  失败: {f} — {e}')


if __name__ == '__main__':
    convert_all()
